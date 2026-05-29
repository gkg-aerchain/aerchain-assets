"""Shared layout engine: one element spec -> HTML preview AND PPTX.
Coordinates in inches. Slide = 20 x 11.25 in (1920x1080 @ 96dpi)."""
import os, html as _html
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

PXIN = 96.0
EMU_IN = 914400
ASSETS = os.path.join(os.path.dirname(__file__), "assets")

# ---- design tokens ----
C = dict(
    purple="7C3AED", purple_d="5B21B6", purple_l="E0CFFA", purple_l2="C4A8F7",
    ink="1B1F26", body="484E59", muted="6A707B", hair="E2E4EA",
    bg="FAFBFC", white="FFFFFF", navy="161A2E",
    green="146B3F", green_bg="E3F6EB", amber="B45309",
    unilever="0B4DA2",
)

FONT_PPTX = {"mont":"Montserrat","mont-med":"Montserrat Medium","mont-light":"Montserrat Light",
             "mont-semi":"Montserrat SemiBold","mont-bold":"Montserrat","mono":"Consolas"}
FONT_W = {"mont":400,"mont-med":500,"mont-light":300,"mont-semi":600,"mont-bold":700,"mono":600}

def R(text, size, color, font="mont", bold=False, italic=False, spc=None):
    return dict(t=text, s=size, c=color, f=font, b=bold, i=italic, spc=spc)

def grad(c0, c1, angle=135):
    """Linear-gradient fill spec (CSS-style angle in degrees)."""
    return dict(g=[c0, c1], angle=angle)

class Slide:
    def __init__(self):
        self.els=[]
    def rect(self,x,y,w,h,fill=None,line=None,lw=0.75,radius=0.0):
        self.els.append(dict(k="rect",x=x,y=y,w=w,h=h,fill=fill,line=line,lw=lw,radius=radius)); return self
    def line(self,x,y,w,h,color,lw=0.75):
        self.els.append(dict(k="line",x=x,y=y,w=w,h=h,color=color,lw=lw)); return self
    def text(self,x,y,w,h,runs,align="l",valign="t",ls=1.0,sa=0.0,wrap=True):
        if isinstance(runs,dict): runs=[[runs]]
        if runs and isinstance(runs[0],dict): runs=[runs]  # single paragraph
        self.els.append(dict(k="text",x=x,y=y,w=w,h=h,runs=runs,align=align,valign=valign,ls=ls,sa=sa,wrap=wrap)); return self
    def img(self,x,y,w,h,path):
        self.els.append(dict(k="img",x=x,y=y,w=w,h=h,path=path)); return self

# ---------------- HTML renderer ----------------
def _runs_html(para):
    out=[]
    for r in para:
        st=f"font-size:{r['s']*1.333:.1f}px;color:#{r['c']};font-weight:{700 if r['b'] else FONT_W[r['f']]};"
        fam="'Montserrat',sans-serif" if r['f']!="mono" else "'JetBrains Mono','Consolas',monospace"
        st+=f"font-family:{fam};"
        if r['i']: st+="font-style:italic;"
        if r['spc'] is not None: st+=f"letter-spacing:{r['spc']}px;"
        if r['f']=="mono": st+="letter-spacing:0.5px;"
        out.append(f"<span style=\"{st}\">{_html.escape(r['t'])}</span>")
    return "".join(out)

def slide_html(sl):
    parts=[]
    for e in sl.els:
        x,y=e['x']*PXIN,e['y']*PXIN
        if e['k']=="rect":
            w,h=e['w']*PXIN,e['h']*PXIN
            st=f"position:absolute;left:{x}px;top:{y}px;width:{w}px;height:{h}px;"
            if isinstance(e['fill'],dict):
                st+=f"background:linear-gradient({e['fill']['angle']}deg,#{e['fill']['g'][0]},#{e['fill']['g'][1]});"
            elif e['fill']: st+=f"background:#{e['fill']};"
            if e['line']: st+=f"border:{e['lw']*1.333:.2f}px solid #{e['line']};"
            if e['radius']: st+=f"border-radius:{e['radius']*PXIN}px;"
            parts.append(f"<div style='{st}'></div>")
        elif e['k']=="line":
            w,h=e['w']*PXIN,e['h']*PXIN
            if e['h']==0:
                st=f"position:absolute;left:{x}px;top:{y}px;width:{w}px;height:0;border-top:{e['lw']*1.333:.2f}px solid #{e['color']};"
            else:
                st=f"position:absolute;left:{x}px;top:{y}px;height:{h}px;width:0;border-left:{e['lw']*1.333:.2f}px solid #{e['color']};"
            parts.append(f"<div style='{st}'></div>")
        elif e['k']=="text":
            w,h=e['w']*PXIN,e['h']*PXIN
            ai={"l":"flex-start","c":"center","r":"flex-end"}[e['align']]
            va={"t":"flex-start","m":"center","b":"flex-end"}[e['valign']]
            ta={"l":"left","c":"center","r":"right"}[e['align']]
            st=(f"position:absolute;left:{x}px;top:{y}px;width:{w}px;height:{h}px;display:flex;flex-direction:column;"
                f"justify-content:{va};align-items:{ai};text-align:{ta};line-height:{e['ls']};overflow:visible;")
            inner=""
            for p in e['runs']:
                inner+=f"<div style='margin-bottom:{e['sa']*1.333:.1f}px;width:100%;text-align:{ta};'>{_runs_html(p)}</div>"
            parts.append(f"<div style='{st}'>{inner}</div>")
        elif e['k']=="img":
            w,h=e['w']*PXIN,e['h']*PXIN
            ap=os.path.abspath(e['path'])
            parts.append(f"<img src='file://{ap}' style='position:absolute;left:{x}px;top:{y}px;width:{w}px;height:{h}px;object-fit:contain;'/>")
    return f"<div class='slide'>{''.join(parts)}</div>"

def write_html(slides, path, title="Preview"):
    body="\n".join(slide_html(s) for s in slides)
    doc=f"""<!doctype html><html><head><meta charset='utf-8'>
<link rel='preconnect' href='https://fonts.googleapis.com'>
<link href='https://fonts.googleapis.com/css2?family=Montserrat:wght@300;400;500;600;700;800&family=JetBrains+Mono:wght@400;500;600&display=swap' rel='stylesheet'>
<style>*{{margin:0;padding:0;box-sizing:border-box}}body{{background:#333}}
.slide{{position:relative;width:1920px;height:1080px;background:#fff;overflow:hidden;margin:0 auto 24px}}</style>
</head><body>{body}</body></html>"""
    with open(path,"w") as f: f.write(doc)

# ---------------- PPTX renderer ----------------
def _set_spc(run, spc_pt):
    rPr=run._r.get_or_add_rPr(); rPr.set('spc', str(int(spc_pt*100)))

def render_pptx(slides, path):
    prs=Presentation(); prs.slide_width=Inches(20); prs.slide_height=Inches(11.25)
    blank=prs.slide_layouts[6]
    for sl in slides:
        s=prs.slides.add_slide(blank)
        for e in sl.els:
            if e['k']=="rect":
                shp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if e['radius'] else MSO_SHAPE.RECTANGLE,
                    Inches(e['x']),Inches(e['y']),Inches(e['w']),Inches(e['h']))
                if e['radius']:
                    try: shp.adjustments[0]=min(0.5, e['radius']/min(e['w'],e['h']))
                    except: pass
                if isinstance(e['fill'],dict):
                    shp.fill.gradient()
                    stops=shp.fill.gradient_stops
                    stops[0].position=0.0; stops[0].color.rgb=RGBColor.from_string(e['fill']['g'][0])
                    stops[1].position=1.0; stops[1].color.rgb=RGBColor.from_string(e['fill']['g'][1])
                    try: shp.fill.gradient_angle=(e['fill']['angle']-90)%360
                    except Exception: pass
                elif e['fill']:
                    shp.fill.solid(); shp.fill.fore_color.rgb=RGBColor.from_string(e['fill'])
                else:
                    shp.fill.background()
                if e['line']:
                    shp.line.color.rgb=RGBColor.from_string(e['line']); shp.line.width=Pt(e['lw'])
                else:
                    shp.line.fill.background()
                shp.shadow.inherit=False
            elif e['k']=="line":
                x2=Inches(e['x']+e['w']); y2=Inches(e['y']+e['h'])
                cn=s.shapes.add_connector(2, Inches(e['x']),Inches(e['y']),x2,y2)
                cn.line.color.rgb=RGBColor.from_string(e['color']); cn.line.width=Pt(e['lw'])
                cn.shadow.inherit=False
            elif e['k']=="text":
                tb=s.shapes.add_textbox(Inches(e['x']),Inches(e['y']),Inches(e['w']),Inches(e['h']))
                tf=tb.text_frame; tf.word_wrap=e['wrap']
                from pptx.enum.text import MSO_AUTO_SIZE
                tf.auto_size=MSO_AUTO_SIZE.NONE
                for m in ("margin_left","margin_right","margin_top","margin_bottom"):
                    setattr(tf,m,0)
                tf.vertical_anchor={"t":MSO_ANCHOR.TOP,"m":MSO_ANCHOR.MIDDLE,"b":MSO_ANCHOR.BOTTOM}[e['valign']]
                for pi,para in enumerate(e['runs']):
                    p=tf.paragraphs[0] if pi==0 else tf.add_paragraph()
                    p.alignment={"l":PP_ALIGN.LEFT,"c":PP_ALIGN.CENTER,"r":PP_ALIGN.RIGHT}[e['align']]
                    p.line_spacing=e['ls']; p.space_after=Pt(e['sa']); p.space_before=Pt(0)
                    for r in para:
                        run=p.add_run(); run.text=r['t']
                        run.font.name=FONT_PPTX[r['f']]; run.font.size=Pt(r['s'])
                        run.font.bold=bool(r['b']); run.font.italic=bool(r['i'])
                        run.font.color.rgb=RGBColor.from_string(r['c'])
                        if r['spc'] is not None: _set_spc(run, r['spc'])
            elif e['k']=="img":
                s.shapes.add_picture(e['path'],Inches(e['x']),Inches(e['y']),Inches(e['w']),Inches(e['h']))
    prs.save(path)

def render_png(slides, html_path, out_prefix, scale=1.0):
    """Render each slide via Chromium to PNG."""
    write_html(slides, html_path)
    from playwright.sync_api import sync_playwright
    paths=[]
    with sync_playwright() as p:
        b=p.chromium.launch()
        pg=b.new_page(viewport={"width":1920,"height":1080},device_scale_factor=scale)
        pg.goto("file://"+os.path.abspath(html_path))
        pg.wait_for_timeout(1200)
        els=pg.query_selector_all(".slide")
        for i,el in enumerate(els):
            fp=f"{out_prefix}_{i+1}.png"; el.screenshot(path=fp); paths.append(fp)
        b.close()
    return paths
